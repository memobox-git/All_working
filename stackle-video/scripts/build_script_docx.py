"""Turns a Stackle script (Markdown) into a plain Word doc for reading.

Usage: python3 build_script_docx.py senior-to-staff-de-notion-walkthrough.md
Writes the .docx next to the .md. Spoken lines are body text; cues in brackets
become short grey lines so the words read straight through.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

src = Path(sys.argv[1])
md = src.read_text()
script = md[md.index("## Script"):md.index("## Notes for this version")]
minutes = re.search(r"\*\*Estimated length:\*\* ([^\n]+)", md).group(1).strip()
GREY = RGBColor(0x6B, 0x70, 0x80)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.25

doc.add_heading("Senior to Staff Data Engineer: My 15-Week Plan", level=0)
meta = doc.add_paragraph(f"YouTube script, Notion screen-share version. Length: {minutes}.")
meta.runs[0].font.color.rgb = GREY


def add_text(p, s):
    # Keep **bold** from the script; drop the other Markdown marks.
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", s)):
        p.add_run(part.replace("*", "")).bold = i % 2 == 1


def cue_text(c):
    c = c.rstrip(":").strip()
    c = re.sub(r",? full screen", "", c)
    c = re.sub(r"(Graphic \d+|Card \d+): \S+", r"\1", c)
    return c[0].upper() + c[1:].lower() if c.isupper() else c


for block in re.split(r"\n\s*\n", script.split("\n", 1)[1]):
    b = block.strip()
    if not b or b == "---":
        continue
    if b.startswith("### "):
        doc.add_heading(b[4:], level=1)
        continue
    if b.startswith("#### "):
        doc.add_heading(b[5:], level=2)
        continue
    m = re.match(r"^\*\*((?:\[[^\]]+\]\s*)+)\*\*\s*(.*)$", b, re.S)
    if m:
        cue = doc.add_paragraph()
        run = cue.add_run("[" + "]  [".join(cue_text(c) for c in re.findall(r"\[([^\]]+)\]", m.group(1))) + "]")
        run.italic, run.font.size, run.font.color.rgb = True, Pt(10), GREY
        cue.paragraph_format.space_after = Pt(2)
        if m.group(2).strip():
            add_text(doc.add_paragraph(), m.group(2).strip())
        continue
    add_text(doc.add_paragraph(), b)

out = src.with_suffix(".docx")
doc.save(out)
print("wrote", out)
